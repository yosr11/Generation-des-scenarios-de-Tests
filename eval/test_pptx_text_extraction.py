from pathlib import Path
import sys

from pptx import Presentation


def read_pptx_text(pptx_path: Path) -> str:
    slides = []
    presentation = Presentation(pptx_path)
    for index, slide in enumerate(presentation.slides, start=1):
        texts = []
        for shape in slide.shapes:
            if not shape.has_text_frame:
                continue
            for paragraph in shape.text_frame.paragraphs:
                paragraph_text = paragraph.text.strip()
                if paragraph_text:
                    texts.append(paragraph_text)
        if texts:
            slides.append(f"--- Slide {index} ---\n" + "\n".join(texts))
    return "\n\n".join(slides)


def main() -> None:
    if len(sys.argv) != 2:
        print(f"Usage: {sys.argv[0]} <file.pptx>")
        raise SystemExit(1)

    path = Path(sys.argv[1])
    if not path.exists():
        raise SystemExit(f"File not found: {path}")

    print(read_pptx_text(path))


if __name__ == "__main__":
    main()
