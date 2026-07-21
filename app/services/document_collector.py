"""
app/services/document_collector.py
───────────────────────────────────
Collecte les documents liés à une story Jira pour le RAG :
  1. Pièces jointes (PJ) de la story
  2. Pièces jointes de l'epic parent (RAG epic uniquement)
  3. Descriptions des tickets liés (depuis issuelinks en base, sans PJ liées)

Chaque document collecté est retourné sous forme de dict :
  {
    "source": "attachment" | "linked_issue" | "epic_attachment",
    "origin_key": "NUXEPM-404",
    "filename": "SFD_Conges.pdf",
    "text": "... contenu extrait ...",
  }
"""

import os
import logging
from typing import Any, Dict, List, Optional

import requests
import urllib3
from dotenv import load_dotenv

from app.utils.document_loader import extract_text_from_bytes
from app.utils.cleaning import clean_text
from app.services.vision_client import describe_image, ocr_only

load_dotenv()
urllib3.disable_warnings(urllib3.exceptions.InsecureRequestWarning)

logger = logging.getLogger(__name__)

JIRA_BASE_URL = os.getenv("JIRA_BASE_URL", "https://hra-jira.ptx.fr.sopra").rstrip("/")
JIRA_USERNAME = os.getenv("JIRA_USERNAME")
JIRA_PASSWORD = os.getenv("JIRA_PASSWORD")
JIRA_EPIC_LINK_FIELD = os.getenv("JIRA_EPIC_LINK_FIELD", "customfield_12402")

# Taille max d'un fichier à télécharger (10 Mo)
MAX_FILE_SIZE = 10 * 1024 * 1024

# Limite du nombre d'images traitées par VLM par story (réduit les coûts)
MAX_VLM_IMAGES_PER_STORY = int(os.getenv("MAX_VLM_IMAGES_PER_STORY", "9999"))

# PPTX extraction settings
PPTX_VLM_ENABLED = os.getenv("PPTX_VLM_ENABLED", "true").lower() == "true"
PPTX_MAX_VLM_IMAGES = int(os.getenv("PPTX_MAX_VLM_IMAGES", "3"))
PPTX_MIN_IMAGE_SIZE = int(os.getenv("PPTX_MIN_IMAGE_SIZE", "50000"))  # 50KB
PPTX_VLM_PRIORITY_SLIDES = int(os.getenv("PPTX_VLM_PRIORITY_SLIDES", "3"))

# Formats ignorés
_SKIP_EXTENSIONS = (".svg", ".ico")

_session = requests.Session()
_session.auth = (JIRA_USERNAME, JIRA_PASSWORD)
_session.verify = False
_session.headers.update({"Accept": "application/json"})


# ══════════════════════════════════════════════════════════════
#  API PUBLIQUE
# ══════════════════════════════════════════════════════════════


def collect_story_attachments_only(
    issue_key: str,
    *,
    vlm_enabled: bool = True,
) -> List[Dict[str, Any]]:
    """
    PJ directement attachées à la story (pas epic, pas tickets liés).
    Utilisé par Agent 1 pour injecter la spec en direct.
    """
    documents: List[Dict[str, Any]] = []
    story_attachments = _get_attachments(issue_key)
    vlm_count = 0

    for att in story_attachments:
        fname = att.get("filename", "") or ""
        name_lower = fname.lower()
        is_image = name_lower.endswith(
            (".png", ".jpg", ".jpeg", ".gif", ".bmp", ".webp")
        )

        text = None
        if is_image:
            if vlm_enabled and vlm_count < MAX_VLM_IMAGES_PER_STORY:
                text = _download_and_extract(att, vlm_enabled=True)
                if text:
                    vlm_count += 1
            elif is_image:
                try:
                    resp = _session.get(att.get("content"), timeout=30)
                    if resp.status_code == 200:
                        text = ocr_only(resp.content, fname)
                except Exception:
                    text = None
        else:
            text = _download_and_extract(att, vlm_enabled=False)

        if text:
            documents.append(
                {
                    "source": "attachment",
                    "origin_key": issue_key,
                    "filename": fname,
                    "text": text,
                }
            )

    logger.info(
        "Collecte PJ story %s : %d document(s) (%d PJ listées)",
        issue_key,
        len(documents),
        len(story_attachments),
    )
    return documents


def linked_docs_from_issuelinks(
    issuelinks: List[Dict[str, Any]],
) -> List[Dict[str, Any]]:
    """
    Construit des documents texte à partir des issuelinks déjà enrichis
    (type, direction, key, status, summary, description) — sans appel Jira.
    """
    documents: List[Dict[str, Any]] = []
    seen_keys: set[str] = set()

    for link in issuelinks or []:
        if not isinstance(link, dict):
            continue
        key = (link.get("key") or "").strip()
        if not key or key in seen_keys:
            continue
        seen_keys.add(key)

        parts = []
        summary = (link.get("summary") or "").strip()
        if summary:
            parts.append(f"Résumé : {summary}")
        desc = (link.get("description") or "").strip()
        if not desc:
            desc = (link.get("description_clean") or "").strip()
        if desc and len(desc) > 20:
            parts.append(desc)
        elif summary:
            parts.append(summary)

        if not parts:
            continue

        text = clean_text("\n".join(parts))
        if len(text.strip()) <= 20:
            continue

        documents.append(
            {
                "source": "linked_issue",
                "origin_key": key,
                "filename": f"{key}_description",
                "text": text,
            }
        )

    return documents


def collect_documents_for_story(
    issue_key: str,
    *,
    issuelinks: Optional[List[Dict[str, Any]]] = None,
    include_epic_attachments: bool = False,
    include_linked: bool = True,
) -> List[Dict[str, Any]]:
    """
    Collecte les documents liés à une story.
    Par défaut : PJ story + descriptions tickets liés (depuis issuelinks, pas de PJ liées).
    """
    documents = collect_story_attachments_only(issue_key, vlm_enabled=True)

    if include_linked:
        links = issuelinks
        if links is None:
            from app.repositories.story_repository import get_story_by_id

            db_story = get_story_by_id(issue_key)
            links = (db_story or {}).get("issuelinks") or []
        documents.extend(linked_docs_from_issuelinks(links))

    if include_epic_attachments:
        epic_key = _get_epic_key_from_db(issue_key)
        if epic_key:
            for att in _get_attachments(epic_key):
                text = _download_and_extract(att, vlm_enabled=False)
                if text:
                    documents.append(
                        {
                            "source": "epic_attachment",
                            "origin_key": epic_key,
                            "filename": att.get("filename", ""),
                            "text": text,
                        }
                    )

    logger.info("Collecte documents pour %s : %d documents", issue_key, len(documents))
    return documents


def collect_documents_for_epic(
    epic_key: str,
    *,
    vlm_focus_story: Optional[str] = None,
    cancel_story_id: Optional[str] = None,
) -> Dict[str, Any]:
    """
    Collecte les documents pour les stories d'un epic (cache DB prioritaire).
    Pas de PJ ni fetch Jira pour les tickets liés — descriptions via issuelinks en base.
    """
    from app.repositories.story_repository import get_stories_by_epic_key

    db_stories = get_stories_by_epic_key(epic_key)
    if db_stories:
        stories = [{"id": s["id"], "title": s.get("summary", "")} for s in db_stories]
        stories_by_key = {s["id"]: s for s in db_stories}
        logger.info(
            "Epic %s : %d stories depuis la base locale", epic_key, len(stories)
        )
    else:
        from app.services.epic_service import get_stories_by_epic

        stories = get_stories_by_epic(epic_key)
        stories_by_key = {}

    if not stories:
        return {
            "epic_key": epic_key,
            "total_stories": 0,
            "total_documents": 0,
            "documents_by_story": {},
            "epic_documents": [],
        }

    # PJ de l'epic (une seule fois)
    epic_documents: List[Dict[str, Any]] = []
    for att in _get_attachments(epic_key):
        text = _download_and_extract(att, vlm_enabled=False)
        if text:
            epic_documents.append(
                {
                    "source": "epic_attachment",
                    "origin_key": epic_key,
                    "filename": att.get("filename", ""),
                    "text": text,
                }
            )

    seen_filenames: set = {d["filename"] for d in epic_documents}
    seen_linked_keys: set = set()
    documents_by_story: Dict[str, List[Dict[str, Any]]] = {}

    for story in stories:
        if cancel_story_id:
            from app.utils.pipeline_cancel import check_pipeline_cancelled

            check_pipeline_cancelled(cancel_story_id)

        story_key = story.get("id", "")
        if not story_key:
            continue

        story_docs: List[Dict[str, Any]] = []
        is_focus = vlm_focus_story is not None and story_key == vlm_focus_story
        db_story = stories_by_key.get(story_key)

        # PJ de la story
        story_attachments = _get_attachments(story_key)
        vlm_count = 0
        for att in story_attachments:
            fname = att.get("filename", "") or ""
            dedup_key = f"{story_key}_{fname}"
            if dedup_key in seen_filenames:
                continue
            seen_filenames.add(dedup_key)

            name_lower = fname.lower()
            is_image = name_lower.endswith(
                (".png", ".jpg", ".jpeg", ".gif", ".bmp", ".webp")
            )
            text = None

            if is_image and is_focus and vlm_count < MAX_VLM_IMAGES_PER_STORY:
                text = _download_and_extract(att, vlm_enabled=True)
                if text:
                    vlm_count += 1
            elif is_image:
                try:
                    resp = _session.get(att.get("content"), timeout=30)
                    if resp.status_code == 200:
                        text = ocr_only(resp.content, fname)
                except Exception:
                    text = None
            else:
                text = _download_and_extract(att, vlm_enabled=False)

            if text:
                story_docs.append(
                    {
                        "source": "attachment",
                        "origin_key": story_key,
                        "filename": fname,
                        "text": text,
                    }
                )

        # Tickets liés : description uniquement (issuelinks en base, pas de PJ)
        issuelinks = (db_story or {}).get("issuelinks") or []
        for doc in linked_docs_from_issuelinks(issuelinks):
            linked_key = doc["origin_key"]
            if linked_key in seen_linked_keys:
                continue
            seen_linked_keys.add(linked_key)
            story_docs.append(doc)

        if story_docs:
            documents_by_story[story_key] = story_docs

    total_docs = len(epic_documents) + sum(
        len(docs) for docs in documents_by_story.values()
    )

    logger.info(
        "Collecte epic %s : %d stories, %d documents (epic PJ: %d, vlm_focus=%s)",
        epic_key,
        len(stories),
        total_docs,
        len(epic_documents),
        vlm_focus_story or "aucune",
    )

    return {
        "epic_key": epic_key,
        "total_stories": len(stories),
        "total_documents": total_docs,
        "documents_by_story": documents_by_story,
        "epic_documents": epic_documents,
    }


# ══════════════════════════════════════════════════════════════
#  FONCTIONS INTERNES
# ══════════════════════════════════════════════════════════════


def _should_skip_file(filename: str) -> bool:
    name_lower = (filename or "").lower()
    return name_lower.endswith(_SKIP_EXTENSIONS)


def _get_epic_key_from_db(issue_key: str) -> Optional[str]:
    from app.repositories.story_repository import get_story_by_id

    db_story = get_story_by_id(issue_key)
    if db_story and db_story.get("epic_key"):
        return db_story["epic_key"]
    return _get_epic_key(issue_key)


def _get_attachments(issue_key: str) -> List[Dict[str, Any]]:
    """Récupère la liste des PJ d'un ticket Jira."""
    url = f"{JIRA_BASE_URL}/rest/api/2/issue/{issue_key}"
    try:
        resp = _session.get(url, params={"fields": "attachment"}, timeout=20)
        if resp.status_code != 200:
            return []
        fields = resp.json().get("fields", {}) or {}
        return fields.get("attachment", []) or []
    except requests.RequestException as e:
        logger.warning("Erreur récupération PJ de %s : %s", issue_key, e)
        return []


def _download_and_extract(
    attachment: Dict[str, Any],
    *,
    vlm_enabled: bool = False,
) -> Optional[str]:
    """Télécharge une PJ et extrait le texte."""
    content_url = attachment.get("content")
    filename = attachment.get("filename", "")
    size = attachment.get("size", 0)

    if not content_url:
        return None

    if size > MAX_FILE_SIZE:
        logger.info("Skip %s (trop volumineux : %d octets)", filename, size)
        return None

    if _should_skip_file(filename):
        logger.debug("Skip %s (format ignoré)", filename)
        return None

    try:
        resp = _session.get(content_url, timeout=60)
        if resp.status_code != 200:
            return None
    except requests.RequestException as e:
        logger.warning("Erreur téléchargement %s : %s", filename, e)
        return None

    name_lower = filename.lower()
    if name_lower.endswith((".pptx", ".ppt")):
        return _extract_pptx_with_vlm(resp.content, vlm_enabled=PPTX_VLM_ENABLED)

    is_image = name_lower.endswith((".png", ".jpg", ".jpeg", ".gif", ".bmp", ".webp"))
    if is_image:
        if not vlm_enabled:
            logger.debug("Skip image %s (VLM désactivé pour cette source)", filename)
            return None
        return describe_image(resp.content, filename)

    return extract_text_from_bytes(resp.content, filename)


def _extract_pptx_with_vlm(content: bytes, vlm_enabled: bool = True) -> Optional[str]:
    """PPTX: texte + VLM sur images critiques (slides 1-3, max 3 images, >50KB)."""
    try:
        import io
        from pptx import Presentation
        from pptx.enum.shapes import MSO_SHAPE_TYPE
    except ImportError:
        logger.warning("python-pptx not installed, skipping PPTX")
        return None

    try:
        prs = Presentation(io.BytesIO(content))
        slides_text = []
        image_count = 0

        for slide_num, slide in enumerate(prs.slides, 1):
            slide_text = [f"--- Slide {slide_num} ---"]

            for shape in slide.shapes:
                if hasattr(shape, "text") and shape.text.strip():
                    slide_text.append(shape.text)

                if (
                    vlm_enabled
                    and shape.shape_type == MSO_SHAPE_TYPE.PICTURE
                    and slide_num <= PPTX_VLM_PRIORITY_SLIDES
                    and image_count < PPTX_MAX_VLM_IMAGES
                ):
                    try:
                        image = shape.image
                        image_bytes = image.blob
                        if len(image_bytes) < PPTX_MIN_IMAGE_SIZE:
                            continue
                        desc = describe_image(
                            image_bytes, f"pptx_slide_{slide_num}.png"
                        )
                        if desc:
                            slide_text.append(f"[Image slide {slide_num}: {desc}]")
                            image_count += 1
                            logger.info(
                                f"VLM PPTX: slide {slide_num} ({image_count}/{PPTX_MAX_VLM_IMAGES})"
                            )
                    except Exception as e:
                        logger.debug(f"Erreur VLM image slide {slide_num}: {e}")

            if len(slide_text) > 1:
                slides_text.append("\n".join(slide_text))

        return "\n\n".join(slides_text) if slides_text else None
    except Exception as e:
        logger.warning(f"Erreur extraction PPTX: {e}")
        return None


def _get_epic_key(issue_key: str) -> Optional[str]:
    """Récupère la clé de l'epic parent d'une story (Jira)."""
    from app.services.jira_service import get_epic_for_story

    epic = get_epic_for_story(issue_key)
    return epic["key"] if epic else None
